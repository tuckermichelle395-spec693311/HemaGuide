#!/usr/bin/env python3
"""Generate a balanced, fully synthetic leukemia/lymphoma case set for plots."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = PROJECT_ROOT / "test_data" / "case_similarity_50"


def spec(code, diagnosis, phase, age, sex, ecog, biomarkers, prior, course, question, decision, role="history"):
    return {
        "code": code, "diagnosis": diagnosis, "phase": phase, "age": age,
        "sex": sex, "ecog": ecog, "biomarkers": biomarkers, "prior": prior,
        "course": course, "question": question, "decision": decision, "role": role,
    }


LEUKEMIA = [
    spec("LEU-001", "核心结合因子AML，RUNX1::RUNX1T1阳性", "初诊、适合强化治疗", 41, "女", 0, "t(8;21)，KIT D816V，CD33阳性", "无", "骨髓原始细胞46%，无中枢侵犯，器官功能可。", "请制定诱导、巩固及MRD监测策略。", "采用7+3联合GO诱导，后续高剂量阿糖胞苷巩固，动态监测RUNX1::RUNX1T1 MRD。", "query"),
    spec("LEU-002", "NPM1突变AML", "初诊、中年", 55, "男", 1, "NPM1突变，FLT3-ITD低比值", "无", "白细胞升高，经细胞减灭后稳定，心功能正常。", "是否需联合FLT3抑制剂？", "强化诱导联合FLT3抑制剂，以NPM1作为MRD标志指导巩固和移植评估。"),
    spec("LEU-003", "FLT3-ITD阳性AML", "首次复发", 47, "女", 1, "FLT3-ITD高比值，DNMT3A突变", "7+3联合米哚洛；HiDAC巩固", "缓解14个月后骨髓复发，暂无活动性感染。", "如何选择挽救治疗并衔接移植？", "采用含FLT3抑制剂的挽救方案，获得缓解后尽快进入异基因移植。"),
    spec("LEU-004", "IDH1突变AML", "老年、不适合强化", 76, "男", 2, "IDH1 R132C，SRSF2突变", "无", "合并慢性肾功能不全和反复肺部感染。", "低强度治疗与靶向治疗如何选择？", "结合体能和感染状态选择去甲基化药联合BCL2抑制剂，不耐受时考虑IDH1靶向治疗。"),
    spec("LEU-005", "TP53多命中AML", "难治、不良风险", 63, "女", 1, "复杂核型，TP53 VAF 62%", "CPX-351诱导未缓解", "持续原始细胞增多，有配型同胞供者。", "是否继续挽救并进入移植？", "优先临床试验或非交叉耐药挽救，达到最低疾病负荷后评估移植。"),
    spec("LEU-006", "治疗相关AML", "初诊、高危", 68, "男", 1, "复杂核型，ASXL1和RUNX1突变", "乳腺癌化疗史8年", "血细胞减少逐渐加重，骨髓原始细胞31%。", "诱导应采用CPX-351还是常规方案？", "考虑CPX-351诱导，缓解后依据MRD和体能尽快评估异基因移植。"),
    spec("LEU-007", "CBFB::MYH11阳性AML", "巩固后分子复发", 38, "女", 0, "inv(16)，KIT exon 8突变", "7+3联合GO；3个疗程HiDAC", "血液学仍缓解，CBFB::MYH11转录本连续升高。", "分子复发阶段是否应干预？", "重复确认分子学趋势，启动挽救治疗和移植评估，避免等待形态学复发。"),
    spec("LEU-008", "急性早幼粒细胞白血病", "初诊高危、DIC", 34, "男", 2, "PML::RARA阳性，白细胞28×10^9/L", "无", "入院时凝血功能障碍和齿龈出血，已立即启动ATRA。", "如何安排诱导和凝血支持？", "按高危APL立即进行ATRA为基础的诱导，密切纠正凝血异常并防治分化综合征。"),
    spec("LEU-009", "急性早幼粒细胞白血病", "分子复发", 46, "女", 0, "PML::RARA低水平复阳", "ATRA联合ATO获得缓解", "治疗结束18个月后连续两次PCR阳性，无出血。", "分子复发的挽救和巩固策略是什么？", "根据既往暴露选择非交叉耐药的APL挽救方案，获得分子缓解后评估移植巩固。"),
    spec("LEU-010", "Ph阳性B-ALL", "初诊、适合强化", 29, "女", 0, "BCR::ABL1 p190，IKZF1plus", "无", "骨髓原始细胞82%，脑脊液阴性。", "TKI、化疗和免疫治疗应如何组合？", "采用强效TKI联合分阶段ALL治疗，结合BCR::ABL1 MRD决定免疫治疗及移植。"),
    spec("LEU-011", "Ph阴性B-ALL", "第二次复发", 43, "男", 1, "CD19阳性，CD22阳性", "多药化疗；异基因移植；blinatumomab", "移植后18个月骨髓复发，无明显GVHD。", "如何在CAR-T和抗体偶联药之间选择？", "根据抗原表达、既往暴露和可及性选择CD19 CAR-T或CD22导向治疗，并预先规划后续巩固。"),
    spec("LEU-012", "ETP-ALL", "诱导失败", 22, "男", 1, "ETP免疫表型，NOTCH1阴性", "儿童样ALL诱导", "诱导后骨髓原始细胞19%，供者搜索进行中。", "如何挽救并获得移植前缓解？", "采用适合T系的挽救方案或临床试验，尽快降低疾病负荷并衔接异基因移植。"),
    spec("LEU-013", "T-ALL", "巩固期MRD持续阳性", 31, "女", 0, "NOTCH1突变，流式MRD 0.08%", "多药ALL诱导和早期巩固", "形态学缓解，MRD下降不足。", "MRD阳性是否应改变巩固并进入移植？", "复核MRD并进行治疗强化，将持续MRD作为高危因素评估异基因移植。"),
    spec("LEU-014", "混合表型急性白血病", "初诊、系别模糊", 36, "男", 1, "B/髓系MPAL，KMT2A重排", "无", "骨髓同时表达B系与髓系标志，无中枢侵犯。", "应选ALL样还是AML样诱导？", "经专家病理复核后优先采用ALL导向方案，依据MRD尽快评估移植。"),
    spec("LEU-015", "慢性髓性白血病", "慢性期、一线不耐受", 52, "女", 0, "BCR::ABL1 e14a2，无耐药突变", "imatinib后水肿和肌痛", "3个月BCR::ABL1 7.8%，依从性良好。", "如何在不耐受时更换TKI？", "复核交互作用和心血管风险，选择替代TKI并按时监测BCR::ABL1。"),
    spec("LEU-016", "慢性髓性白血病", "加速期、TKI耐药", 61, "男", 1, "BCR::ABL1 T315I", "imatinib；dasatinib；nilotinib", "原始细胞13%，脾大进展，有一名匹配非血缘供者。", "靶向治疗与移植如何衔接？", "使用对T315I有活性的TKI控制疾病，同步完成移植准备。"),
    spec("LEU-017", "慢性髓性白血病急变", "髓系急变", 45, "男", 2, "BCR::ABL1，新增RUNX1突变", "bosutinib治疗3年", "骨髓原始细胞34%，伴肺部感染刚得到控制。", "如何进行急变诱导？", "采用AML样治疗联合合适TKI，获得第二慢性期后尽快移植。"),
    spec("LEU-018", "慢性髓单核细胞白血病", "增殖型、高危", 72, "女", 1, "ASXL1、NRAS和SRSF2突变", "hydroxyurea", "白细胞持续升高，输血依赖，脾大。", "是否应使用去甲基化治疗？", "以症状、细胞减少和进展风险为基础启动去甲基化治疗，体能允许时评估移植。"),
    spec("LEU-019", "慢性髓单核细胞白血病", "向AML转化", 66, "男", 2, "NRAS、SETBP1突变，原始细胞24%", "azacitidine 8个疗程", "进行性细胞减少和盗汗，伴慢性肺病。", "转化后如何选择低强度方案？", "结合器官功能选择AML导向的低强度组合，并评估临床试验。"),
    spec("LEU-020", "高危骨髓增生异常综合征", "进展、输血依赖", 64, "女", 1, "复杂核型，EZH2和ASXL1突变", "支持治疗", "骨髓原始细胞14%，每2周需输红细胞。", "去甲基化治疗后是否应移植？", "启动去甲基化治疗并同步进行供者搜索，疾病稳定后评估移植。"),
    spec("LEU-021", "幼年型髓单核细胞白血病", "新诊断、高危", 7, "男", 1, "PTPN11突变，单体细7", "无", "持续单核细胞增多、脾大和胎儿血红蛋白升高。", "移植前是否需要减瘤治疗？", "尽快启动儿童移植中心评估，根据疾病负荷选择移植前控制策略。"),
    spec("LEU-022", "毛细胞白血病", "症状性初诊", 59, "男", 0, "BRAF V600E阳性", "无", "中性粒细胞减少、反复感染和明显脾大。", "喷司他丁类治疗和BRAF靶向治疗如何选择？", "无活动性严重感染时优先标准核苷类似物治疗，特定情境考虑BRAF导向方案。"),
    spec("LEU-023", "T细胞大颗粒淋巴细胞白血病", "慢性、症状性细胞减少", 69, "女", 1, "STAT3突变，TCR克隆性阳性", "G-CSF间歇支持", "重度中性粒细胞减少和反复口腔感染。", "何时启动免疫抑制治疗？", "对症状性重度细胞减少启动免疫抑制治疗，定期评估血象和感染。"),
    spec("LEU-024", "B细胞前淋巴细胞白血病", "进展性、高肿瘤负荷", 71, "男", 2, "TP53缺失，IGH::MYC阳性", "BTK抑制剂短暂缓解", "白细胞快速上升，脾大和B症状明显。", "多线治疗后有何挽救选择？", "建议专家中心复核病理，优先临床试验或细胞治疗评估。"),
    spec("LEU-025", "系统性肥大细胞增生症伴髓系肿瘤", "进展性、器官损害", 57, "女", 1, "KIT D816V，SRSF2/ASXL1突变", "midostaurin", "脾大、贫血和胆红素升高，药物治疗后再次进展。", "是否更换强效KIT抑制剂并评估移植？", "更换更强效的KIT导向治疗，结合伴发髓系肿瘤和体能评估移植。"),
]


LYMPHOMA = [
    spec("LYM-001", "弥漫性大B细胞淋巴瘤，GCB型", "初诊、晚期", 58, "男", 1, "MYC阴性，BCL2蛋白高表达", "无", "Ann Arbor IV期，多站淋巴结及骨髓受累，IPI 3分。", "一线治疗强度和中枢预防如何安排？", "采用标准免疫化疗，根据CNS-IPI和受累部位个体化评估中枢预防。", "query"),
    spec("LYM-002", "弥漫性大B细胞淋巴瘤，ABC型", "一线难治", 64, "女", 1, "CD19阳性，MYD88 L265P", "R-CHOP 6个疗程", "疗程结束PET-CT显示代谢进展，已重新活检确认。", "二线应直接进入CAR-T还是挽救移植？", "按早期难治DLBCL评估CD19 CAR-T，同步安排桥接治疗和器官功能评估。"),
    spec("LYM-003", "高级别B细胞淋巴瘤", "初诊、双打击", 49, "男", 1, "MYC和BCL2重排", "无", "腹膜后大包块和高LDH，CNS-IPI高危。", "是否采用强化一线并进行CNS预防？", "采用强化免疫化疗，结合中枢风险安排预防并密切监测疗效。"),
    spec("LYM-004", "弥漫性大B细胞淋巴瘤", "CAR-T后复发", 61, "女", 2, "CD19弱阳性，CD20阳性", "R-CHOP；挽救化疗；CD19 CAR-T", "CAR-T后9个月多站复发，伴轻度细胞减少。", "CAR-T后复发如何选择双特异抗体或偶联药？", "根据抗原表达、细胞减少和感染风险选择双特异抗体或抗体偶联药。"),
    spec("LYM-005", "弥漫性大B细胞淋巴瘤", "晚期复发、移植适合", 52, "男", 0, "CD20阳性，无MYC重排", "R-CHOP后缓解36个月", "局部淋巴结复发，已活检，器官功能良好。", "是否采用挽救化疗后自体移植？", "对化疗敏感的晚期复发采用挽救治疗，达缓解后进行自体移植。"),
    spec("LYM-006", "原发中枢神经系淋巴瘤", "初诊、老年", 73, "女", 2, "MYD88 L265P，CD20阳性", "无", "多发脑实质病灶，肾小球滤过率轻度下降。", "高剂量甲氨蝶呤方案如何根据器官功能调整？", "在专科中心采用以高剂量甲氨蝶呤为基础的个体化方案，严密监测肾功能。"),
    spec("LYM-007", "原发纵隔大B细胞淋巴瘤", "一线难治", 35, "女", 1, "CD30阳性，PD-L1高表达", "DA-EPOCH-R 6个疗程；局部放疗", "纵隔病灶持续PET强阳性，活检证实活性肿瘤。", "CAR-T、PD-1抑制剂和移植如何排序？", "评估CD19 CAR-T或临床试验，需桥接时结合PD-1通路特征选择方案。"),
    spec("LYM-008", "原发纵隔大B细胞淋巴瘤", "CAR-T后局部复发", 42, "男", 0, "CD19阴性，CD30阳性", "DA-EPOCH-R；放疗；CD19 CAR-T", "CAR-T后14个月单一纵隔病灶复发。", "局部治疗与全身治疗如何组合？", "经活检和全身分期后考虑局部控制，并根据CD30和PD-1通路选择全身治疗。"),
    spec("LYM-009", "滤泡性淋巴瘤1–2级", "低肿瘤负荷初诊", 60, "女", 0, "t(14;18)，FLIPI 1", "无", "无B症状，结节较小，血象正常。", "是否可以观察等待？", "在无治疗指征时采用观察等待，定期评估症状、血象和肿瘤负荷。"),
    spec("LYM-010", "滤泡性淋巴瘤2级", "高肿瘤负荷初诊", 67, "男", 1, "EZH2突变，FLIPI 3", "无", "巨大腹膜后肿块、脾大和盗汗。", "一线免疫化疗如何选择？", "采用抗CD20为基础的全身治疗，结合疾病负荷、合并症和患者偏好选择联合方案。"),
    spec("LYM-011", "滤泡性淋巴瘤", "POD24、早期复发", 56, "女", 1, "CD20阳性，无转化证据", "bendamustine联合抗CD20抗体", "完成治疗后14个月多站进展，已重新活检。", "POD24后如何选择二线及细胞治疗？", "优先重新活检排除转化，采用非交叉耐药方案并评估双特异抗体或CAR-T。"),
    spec("LYM-012", "套细胞淋巴瘤", "年轻、初诊", 51, "男", 0, "t(11;14)，TP53野生型，ki-67 35%", "无", "IV期伴骨髓受累，器官功能良好。", "是否采用含阿糖胞苷的诱导和自体移植？", "采用含阿糖胞苷和抗CD20的强化诱导，达缓解后评估自体移植及维持治疗。"),
    spec("LYM-013", "套细胞淋巴瘤", "老年初诊", 78, "女", 2, "t(11;14)，ki-67 20%", "无", "巨大脾脏和贫血，合并房颤和慢性肾病。", "如何选择低毒性一线治疗？", "根据心肾功能选择减强度免疫化疗或无化疗策略，并严密管理感染与出血风险。"),
    spec("LYM-014", "套细胞淋巴瘤", "BTK抑制剂后进展", 65, "男", 1, "TP53突变，ki-67 70%", "免疫化疗；自体移植；BTK抑制剂", "疾病快速进展并出现多站节外受累。", "是否应优先CAR-T？", "尽快评估CAR-T或临床试验，用短程桥接治疗控制高增殖性疾病。"),
    spec("LYM-015", "经典型霍奇金淋巴瘤", "早期不良风险", 27, "女", 0, "EBV阴性，纵隔大包块", "无", "Ann Arbor IIB期，ESR升高，肺功能正常。", "化疗疗程数和放疗如何由PET指导？", "采用PET适应性一线方案，根据中期和终末PET决定疗程强度与局部放疗。"),
    spec("LYM-016", "经典型霍奇金淋巴瘤", "晚期初诊", 44, "男", 1, "CD30阳性，IPS 4", "无", "IVB期，骨髓及肺受累，有明显B症状。", "一线应选择BrECADD还是其他方案？", "采用PET适应性的现代强化一线方案，根据年龄和器官功能平衡治愈率与毒性。"),
    spec("LYM-017", "经典型霍奇金淋巴瘤", "自体移植后复发", 39, "女", 1, "CD30阳性，PD-L1高表达", "ABVD；挽救化疗；自体移植", "移植后11个月多站复发，无严重感染。", "应选PD-1抑制剂还是CD30导向治疗？", "根据既往暴露和神经毒性风险选择PD-1或CD30导向治疗，持久缓解后评估后续巩固。"),
    spec("LYM-018", "MALT淋巴瘤", "胃局限期、初诊", 62, "男", 0, "API2::MALT1阴性，幽门螺杆菌阳性", "无", "病灶局限于胃，无深层浸润或远处受累。", "是否首先进行幽门螺杆菌根除？", "首先完成幽门螺杆菌根除，根据内镜和组织学随访决定是否局部放疗。"),
    spec("LYM-019", "脾边缘区淋巴瘤", "症状性、细胞减少", 70, "女", 1, "NOTCH2突变，IGH克隆性", "观察3年", "脾大进展，贫血和血小板减少，无转化证据。", "何时启动抗CD20为基础的治疗？", "症状性脾大及细胞减少已构成治疗指征，采用抗CD20为基础的个体化方案。"),
    spec("LYM-020", "外周T细胞淋巴瘤，NOS", "初诊、高危", 55, "男", 1, "TCR克隆性，CD30阴性", "无", "IV期伴骨髓受累和B症状，PIT高危。", "是否应在首次缓解时自体移植？", "采用T细胞淋巴瘤导向诱导，达缓解后结合风险和疗效评估自体移植。"),
    spec("LYM-021", "血管免疫母T细胞淋巴瘤", "首次复发", 63, "女", 2, "TET2、DNMT3A和RHOA G17V突变", "CHOP样治疗后缓解10个月", "全身淋巴结增大、皮疹和自身免疫性溶血。", "复发时如何选择表观遗传治疗和移植？", "选择非交叉耐药挽救或临床试验，获得缓解后根据体能评估移植。"),
    spec("LYM-022", "ALK阳性间变性大细胞淋巴瘤", "初诊、晚期", 32, "男", 0, "NPM::ALK阳性，CD30强阳性", "无", "III期伴B症状，无中枢受累。", "一线是否应采用CD30导向联合方案？", "采用CD30导向联合化疗，根据PET反应决定后续疗程和巩固。"),
    spec("LYM-023", "伯基特淋巴瘤", "初诊、高肿瘤负荷", 26, "女", 1, "MYC重排，EBV阴性", "无", "腹腔巨大肿块，LDH显著升高，肿瘤溶解风险高。", "如何安排预阶段、强化治疗和CNS预防？", "先进行肿瘤溶解预防和预阶段处理，随后启动短疗程高强度方案并纳入CNS导向治疗。"),
    spec("LYM-024", "鼻型结外NK/T细胞淋巴瘤", "局限期初诊", 48, "男", 1, "EBER阳性，血浆EBV-DNA升高", "无", "鼻腔局部侵袭性病灶，无远处播散。", "放疗与含门冬酰胺酶化疗如何整合？", "采用以放疗为核心的联合策略，整合含门冬酰胺酶方案并监测EBV-DNA。"),
    spec("LYM-025", "移植后淋巴增殖性疾病", "肾移植后、EBV阳性", 50, "女", 2, "EBER阳性，CD20阳性", "肾移植后免疫抑制6年", "多站淋巴结及移植肾周围病灶，肾功能波动。", "如何平衡减少免疫抑制、抗CD20治疗与移植器官风险？", "与移植团队协作减少免疫抑制，启动抗CD20治疗，并密切监测器官功能和病毒负荷。"),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB W3")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E5E7E", 14, 8),
        ("Heading 2", 12.5, "2E5E7E", 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB W3")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_field_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.55)
        cells[1].width = Inches(4.95)
        cells[0].text = label
        cells[1].text = str(value)
        set_cell_shading(cells[0], "E8EEF5")
        for cell in cells:
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.name = "Arial"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB W3")
                    run.font.size = Pt(9.5)
        cells[0].paragraphs[0].runs[0].bold = True
    return table


def make_doc(group, item, output_path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    configure_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run(f"合成病例 {item['code']} | {item['diagnosis']}")
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB W3")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string("173B57")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("完全合成·非真实患者·仅用于HemaGuide相似性检索与绘图测试")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("6B7280")

    doc.add_heading("病例基本信息", level=1)
    add_field_table(doc, [
        ("patient_id", f"SYN-{item['code']}"),
        ("dataset_group", group),
        ("case_role", item["role"]),
        ("tumorboard_type", "血液肿瘤委员会"),
        ("age", f"{item['age']} 岁"),
        ("sex", item["sex"]),
        ("ECOG", item["ecog"]),
        ("main_diagnosis", item["diagnosis"]),
        ("disease_phase", item["phase"]),
    ])

    doc.add_heading("预后与分子因素", level=1)
    doc.add_paragraph(item["biomarkers"])

    doc.add_heading("既往治疗", level=1)
    doc.add_paragraph(item["prior"])

    doc.add_heading("临床经过", level=1)
    doc.add_paragraph(item["course"])

    doc.add_heading("肿瘤委员会问题", level=1)
    doc.add_paragraph(item["question"])

    doc.add_heading("历史会议决定", level=1)
    doc.add_paragraph(item["decision"])

    doc.add_heading("数据使用说明", level=1)
    note = doc.add_paragraph()
    note.add_run("本文档中的所有人口学、诊断、分子、治疗及结局信息均为人工构造，不对应任何真实患者，不得用于临床决策。")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(f"Synthetic HemaGuide similarity dataset | {item['code']}")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string("7A7A7A")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    if OUTPUT_ROOT.exists():
        shutil.rmtree(OUTPUT_ROOT)

    manifest_rows = []
    for group, items in (("leukemia", LEUKEMIA), ("lymphoma", LYMPHOMA)):
        assert len(items) == 25
        for item in items:
            role_dir = "query" if item["role"] == "query" else "history"
            filename = f"SYN_{item['code'].replace('-', '_')}.docx"
            path = OUTPUT_ROOT / group / role_dir / filename
            make_doc(group, item, path)
            manifest_rows.append({
                "case_id": item["code"], "group": group, "role": item["role"],
                "diagnosis": item["diagnosis"], "phase": item["phase"],
                "age": item["age"], "sex": item["sex"], "ECOG": item["ecog"],
                "biomarkers": item["biomarkers"],
                "relative_path": str(path.relative_to(OUTPUT_ROOT)),
            })

    with open(OUTPUT_ROOT / "manifest.csv", "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=manifest_rows[0].keys())
        writer.writeheader()
        writer.writerows(manifest_rows)

    readme = OUTPUT_ROOT / "README.txt"
    readme.write_text(
        "HemaGuide 合成病例相似性测试集\n"
        "===============================\n"
        "共50例：白血病组25例，淋巴瘤组25例。\n"
        "每组包含24例历史病例（history）和1例查询病例（query）。\n"
        "病例覆盖多种亚型、年龄、体能状态、分子标志及治疗阶段。\n"
        "manifest.csv为作图和质量检查用索引。\n"
        "全部内容均为人工合成，不对应任何真实患者，不得用于临床决策。\n",
        encoding="utf-8",
    )
    print(f"Generated {len(manifest_rows)} cases in {OUTPUT_ROOT}")


if __name__ == "__main__":
    main()
