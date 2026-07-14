#!/usr/bin/env python3
"""Generate synthetic DOCX fixtures for HemaGuide MOLECULAR-mode testing."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
QUERY_DIR = ROOT / "query_input"
HISTORY_DIR = ROOT / "kb_input" / "tumorboards"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x66, 0x66, 0x66)


def _set_font(run, size=11, bold=False, color=None):
    run.font.name = "Arial Unicode MS"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def _set_cell_free_document_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    _set_font(header.add_run("HemaGuide | MOLECULAR mode test case"), size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    _set_font(footer.add_run("Synthetic teaching data | No real patient information"), size=8.5, color=MUTED)


def _add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    _set_font(p.add_run(title), size=22, bold=True, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    _set_font(p.add_run(subtitle), size=10.5, color=MUTED)


def _add_field(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    _set_font(p.add_run(f"{label}："), bold=True)
    _set_font(p.add_run(value))


def _add_section(doc, heading, paragraphs):
    doc.add_heading(heading, level=2)
    for text in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        _set_font(p.add_run(text))


def _add_molecular_section(doc, paragraphs):
    # This exact standalone heading is consumed by split_document_at_molecular().
    doc.add_heading("Molekulares Tumorboard", level=1)
    for text in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        _set_font(p.add_run(text))


def _save_case(path, title, subtitle, fields, history, question, molecular):
    doc = Document()
    _set_cell_free_document_styles(doc)
    _add_title(doc, title, subtitle)
    for label, value in fields:
        _add_field(doc, label, value)
    _add_section(doc, "Clinical course", history)
    _add_section(doc, "Tumor board question", [question])
    _add_molecular_section(doc, molecular)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def generate():
    _save_case(
        QUERY_DIR / "SIM_MOL_KIT_query.docx",
        "MOLECULAR query case | KIT D816V",
        "Synthetic fixture for molecular routing, case-retrieval toggles, and evidence output",
        [
            ("Patient ID", "SIM-MOL-QUERY-KIT-001"),
            ("Tumor board type", "Molecular tumor board"),
            ("Age / sex", "42 years / male"),
            ("ECOG", "1"),
            ("Main diagnosis", "Core-binding-factor AML, RUNX1::RUNX1T1 positive"),
            ("Current status", "New diagnosis; no prior AML-specific induction therapy"),
        ],
        [
            "The patient presented with fatigue, fever, and thrombocytopenia. Bone marrow blasts were 61%, and flow cytometry supported AML.",
            "Cytogenetics showed t(8;21)(q22;q22.1), with quantitative RUNX1::RUNX1T1 positivity. Cardiac, hepatic, and renal function permitted intensive therapy; infection was controlled.",
        ],
        "Assess the oncogenicity of KIT p.D816V, its implications for relapse risk and MRD management, and propose treatment using historical cases, literature, and meeting evidence.",
        [
            "NGS variant: KIT (NM_000222.3) c.2447A>T, p.Asp816Val (p.D816V), VAF 37.2%, COSMIC ID COSM1314; laboratory classification: Oncogenic.",
            "Cytogenetics/FISH: t(8;21)(q22;q22.1)/RUNX1::RUNX1T1 positive in 82% of evaluated cells.",
            "Molecular tumor board decision: determine whether KIT D816V modifies treatment, consolidation, or MRD follow-up in core-binding-factor AML.",
        ],
    )

    histories = [
        (
            "SIM_MOL_KIT_history_01.docx",
            "History case 01 | High relevance KIT D816V",
            "SIM-MOL-HIST-KIT-001",
            "36 years / female",
            "Core-binding-factor AML, RUNX1::RUNX1T1 positive",
            [
                "At diagnosis, bone marrow blasts were 48%. The patient achieved complete remission after 7+3 plus gemtuzumab ozogamicin induction.",
                "After two high-dose cytarabine consolidation cycles, RUNX1::RUNX1T1 transcripts declined continuously; hematologic remission persisted at 24 months.",
            ],
            "Historical decision: intensive induction plus GO followed by high-dose cytarabine; KIT alone did not determine transplant in first remission, with close molecular MRD monitoring.",
            [
                "NGS variant: KIT (NM_000222.3) c.2447A>T, p.Asp816Val (p.D816V), VAF 31%, COSMIC ID COSM1314; classification: Oncogenic.",
                "FISH/fusion: t(8;21)/RUNX1::RUNX1T1 positive in 76% of evaluated cells.",
                "Outcome: complete remission after induction and molecular MRD negativity after consolidation.",
            ],
        ),
        (
            "SIM_MOL_KIT_history_02.docx",
            "History case 02 | KIT D816V with persistent MRD",
            "SIM-MOL-HIST-KIT-002",
            "57 years / male",
            "Core-binding-factor AML, RUNX1::RUNX1T1 positive",
            [
                "Standard intensive induction produced morphologic complete remission, but RUNX1::RUNX1T1 MRD reduction was inadequate.",
                "Molecular MRD remained positive after two consolidation cycles. Multidisciplinary review led to allogeneic transplant evaluation, modeling an MRD-driven escalation decision.",
            ],
            "Historical decision: KIT alone did not determine transplant; persistent molecular MRD prompted intensified salvage and allogeneic transplant assessment.",
            [
                "NGS variant: KIT (NM_000222.3) c.2447A>T, p.Asp816Val (p.D816V), VAF 42%, COSMIC ID COSM1314; classification: Oncogenic.",
                "FISH/fusion: t(8;21)/RUNX1::RUNX1T1 positive in 69% of evaluated cells.",
                "Outcome: morphologic remission with persistent RUNX1::RUNX1T1 molecular MRD positivity.",
            ],
        ),
        (
            "SIM_MOL_KIT_history_03.docx",
            "History case 03 | Same KIT gene, different variant",
            "SIM-MOL-HIST-KIT-003",
            "49 years / female",
            "Core-binding-factor AML, CBFB::MYH11 positive",
            [
                "The patient achieved complete remission after intensive induction and cytarabine consolidation. The KIT variant differed from the query case.",
                "Response was followed using quantitative CBFB::MYH11 and flow MRD; no relapse occurred during 18 months of follow-up.",
            ],
            "Historical decision: treat as core-binding-factor AML; record KIT p.N822K as risk context, while escalation depends on MRD kinetics rather than gene-name matching alone.",
            [
                "NGS variant: KIT (NM_000222.3) c.2466T>G, p.Asn822Lys (p.N822K), VAF 18%; laboratory classification: Likely Oncogenic.",
                "FISH/fusion: inv(16)/CBFB::MYH11 positive in 71% of evaluated cells.",
                "Outcome: complete remission with CBFB::MYH11 molecular MRD negativity.",
            ],
        ),
        (
            "SIM_MOL_CONTROL_history_04.docx",
            "History case 04 | Different-gene negative control",
            "SIM-MOL-HIST-CTRL-004",
            "51 years / male",
            "NPM1-mutated AML, FLT3-ITD positive",
            [
                "The patient received intensive induction plus a FLT3 inhibitor, achieved complete remission, and proceeded to consolidation.",
                "This case contains no KIT alteration and tests that molecular gene matching does not select a case solely because the disease name is AML.",
            ],
            "Historical decision: use FLT3-targeted therapy and monitor NPM1 MRD; the strategy has no direct correspondence to KIT D816V.",
            [
                "NGS variant: FLT3-ITD, VAF 24%; laboratory classification: Oncogenic.",
                "NGS variant: NPM1 (NM_002520.7) c.860_863dup, p.Trp288CysfsTer12, VAF 39%; classification: Oncogenic.",
                "FISH/karyotype: normal karyotype; common AML rearrangement probes were negative.",
            ],
        ),
    ]

    for filename, title, patient_id, age_sex, diagnosis, history, decision, molecular in histories:
        _save_case(
            HISTORY_DIR / filename,
            title,
            "Synthetic history fixture for gene matching, relevance filtering, and case-retrieval toggles",
            [
                ("Patient ID", patient_id),
                ("Tumor board type", "Molecular tumor board"),
                ("Age / sex", age_sex),
                ("ECOG", "0–1"),
                ("Main diagnosis", diagnosis),
                ("Data status", "Fully synthetic; no real patient correspondence"),
            ],
            history,
            decision,
            molecular,
        )

    print("Generated 1 molecular query and 4 molecular history cases.")


if __name__ == "__main__":
    generate()
